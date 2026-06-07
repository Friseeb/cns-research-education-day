"""
Import submissions (with training_level), judges, and judge assignments from
the 'Judge Assisgnment.xlsx' format (Sheet1).

Usage:
    python manage.py import_judge_assignment_xlsx <path_to_xlsx>

The spreadsheet has:
  col A  Judges        – judge names for the session/tour (only on first row of block)
  col B  First Name
  col C  Last Name
  col D  E-mail Address (presenter)
  col E  I am currently (training status)
  col F  Topic (abstract title)
  col G  First Author
  col H  Co-Author
  col I  Supervisor
  col J  Oral or Poster
  col K  Tag  (PLAT-1, POST-1, …) → abstract_number
  col L  Session
  col M  Start time
  col N  Room
"""

import re

import openpyxl
from django.core.management.base import BaseCommand, CommandError

from judging.models import Category, Event, Judge, JudgeAssignment, PresentationFormat, Submission
from judging.services.imports import _normalize_training_level

_SESSION_PREFIX = re.compile(r"Poster Tour \w+:\s*")


def _parse_judge_names(cell_value):
    """Return a list of judge names from a cell that may contain multiple names."""
    raw = _SESSION_PREFIX.sub("", str(cell_value))
    names = [n.strip() for n in re.split(r"[/\n\r]", raw) if n.strip()]
    return names


def _first_value(cell_value):
    """If a cell contains multiple values (newline-separated), return just the first."""
    if cell_value is None:
        return ""
    return str(cell_value).split("\n")[0].split("\r")[0].strip()


def _judge_email(name):
    """Generate a placeholder email from a judge name."""
    slug = re.sub(r"[^a-z]", "", name.lower().replace(" ", "."))
    return f"{slug}@judges.cns2026.placeholder"


class Command(BaseCommand):
    help = "Import judges and assignments from the CNS Judge Assignment xlsx"

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str, help="Path to the Judge Assignment xlsx file")
        parser.add_argument("--event", type=str, default="CNS Research Day 2026")
        parser.add_argument(
            "--no-submissions",
            action="store_true",
            help="Skip submission upsert (only create judges and assignments)",
        )

    def handle(self, *args, **options):
        try:
            wb = openpyxl.load_workbook(options["xlsx_path"])
        except FileNotFoundError:
            raise CommandError(f"File not found: {options['xlsx_path']}")

        from datetime import date

        event, _ = Event.objects.get_or_create(
            name=options["event"],
            defaults={"is_active": True, "date": date(2026, 6, 5)},
        )
        self.stdout.write(f"Event: {event.name}")

        oral_format, _ = PresentationFormat.objects.get_or_create(name="Oral presentation")
        poster_format, _ = PresentationFormat.objects.get_or_create(name="Poster presentation")

        ws = wb["Sheet1"]
        rows = list(ws.iter_rows(values_only=True))

        # Prepass: collect all unique judge names so we can create them before assigning.
        all_judge_names = set()
        for row in rows:
            judges_cell = row[0]
            if judges_cell and str(judges_cell) not in ("Judges", "Tag"):
                for name in _parse_judge_names(judges_cell):
                    if name:
                        all_judge_names.add(name)

        self.stdout.write(f"Judges found in file: {sorted(all_judge_names)}")

        # Create/get Judge objects.
        judge_objects = {}
        judges_created = 0
        for name in sorted(all_judge_names):
            email = _judge_email(name)
            judge, created = Judge.objects.get_or_create(
                event=event,
                email=email,
                defaults={"name": name},
            )
            if created:
                judges_created += 1
            judge_objects[name] = judge

        self.stdout.write(f"Judges created: {judges_created}  (total: {len(judge_objects)})")

        # Main pass: update submissions and create assignments.
        current_judge_names = []
        subs_created = 0
        subs_updated = 0
        assignments_created = 0
        assignments_skipped = 0

        for row in rows:
            judges_cell = row[0]
            first_name_raw = row[1]
            last_name_raw = row[2]
            training_raw = str(row[4] or "").strip()
            title = str(row[5] or "").strip()
            first_author = str(row[6] or "").strip()
            co_authors = str(row[7] or "").strip()
            format_raw = str(row[9] or "").strip().lower()
            tag = str(row[10] or "").strip()
            session = str(row[11] or "").strip()
            time_str = str(row[12] or "").strip()
            room = str(row[13] or "").strip()

            # Update current judges whenever the cell is populated.
            if judges_cell and str(judges_cell) not in ("Judges", "Tag"):
                current_judge_names = _parse_judge_names(judges_cell)

            # Skip header rows and blank rows.
            if not tag or tag in ("Tag", "Judges") or not first_name_raw:
                continue

            first_name = _first_value(first_name_raw)
            last_name = _first_value(last_name_raw)
            presenting_author = f"{first_name} {last_name}".strip()

            training_level = _normalize_training_level(training_raw)
            fmt = oral_format if "oral" in format_raw else poster_format
            location_full = " | ".join(filter(None, [session, time_str, room]))

            if not options["no_submissions"]:
                # Determine category from training role.
                category_name = _category_name(training_raw)
                category, _ = Category.objects.get_or_create(event=event, name=category_name)

                sub, created_flag = Submission.objects.update_or_create(
                    event=event,
                    abstract_number=tag,
                    defaults={
                        "title": title,
                        "presenting_author": presenting_author,
                        "co_authors": co_authors,
                        "category": category,
                        "presentation_format": fmt,
                        "training_level": training_level,
                        "abstract_text": title,
                        "location": location_full,
                        "is_active": True,
                    },
                )
                if created_flag:
                    subs_created += 1
                else:
                    subs_updated += 1

            # Create judge assignments.
            submission = Submission.objects.filter(event=event, abstract_number=tag).first()
            if not submission:
                continue

            for judge_name in current_judge_names:
                judge = judge_objects.get(judge_name)
                if not judge:
                    continue
                _, created_flag = JudgeAssignment.objects.get_or_create(
                    event=event,
                    judge=judge,
                    submission=submission,
                )
                if created_flag:
                    assignments_created += 1
                else:
                    assignments_skipped += 1

        if not options["no_submissions"]:
            self.stdout.write(
                f"Submissions: {subs_created} created, {subs_updated} updated"
            )
        self.stdout.write(
            f"Assignments: {assignments_created} created, {assignments_skipped} already existed"
        )
        self.stdout.write(self.style.SUCCESS("Done."))
        self.stdout.write(
            "\nNOTE: Judge emails are placeholders. Update them in Django admin\n"
            "before generating login links.\n"
        )


def _category_name(role_raw):
    r = role_raw.lower()
    if "resident" in r:
        return "Resident"
    if any(x in r for x in ("fellow", "post-doc", "postdoc", "associate", "assistant")):
        return "Fellow / Research Staff"
    if any(x in r for x in ("medical student", "undergraduate", "master", "phd", "graduate")):
        return "Student"
    return role_raw or "Other"
